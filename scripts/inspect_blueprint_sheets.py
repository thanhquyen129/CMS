import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('docs/po/21102026-LCMS_BA_DataEntry_TO_BE_v0.9_Package/LCMS_TO_BE_Field_Optimization_Inheritance_Matrix_v0.9.xlsx', data_only=True)

for sheet_name in ['03 DEV Backlog', '04 UAT Cases', '05 GAP và rủi ro', '06 Quy tắc kế thừa']:
    print(f"\n=================== {sheet_name} ===================")
    ws = wb[sheet_name]
    for r in range(1, ws.max_row + 1):
        vals = [str(ws.cell(r, c).value or '').strip().replace('\n', ' ') for c in range(1, ws.max_column + 1)]
        if any(vals):
            print(f"[{r:02d}] " + " | ".join(vals))
