import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

md_path = "docs/po/22102026/LCMS_BA_Final_Verification_DEV_Response_20261022.md"
docx_path = "docs/po/22102026/LCMS_BA_Final_Verification_DEV_Response_20261022.docx"

doc = docx.Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Colors
C_PRIMARY = RGBColor(24, 76, 120)     # Deep Blue
C_SECONDARY = RGBColor(41, 128, 185)  # Bright Blue
C_TEXT = RGBColor(44, 62, 80)         # Charcoal Text
C_MUTED = RGBColor(127, 140, 141)     # Muted Gray

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

with open(md_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

in_table = False
table_data = []

def flush_table():
    global in_table, table_data
    if not table_data:
        in_table = False
        return
    
    num_cols = max(len(row) for row in table_data)
    table = doc.add_table(rows=len(table_data), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for r_idx, row_cells in enumerate(table_data):
        is_header = (r_idx == 0)
        for c_idx, text in enumerate(row_cells):
            if c_idx >= num_cols:
                continue
            cell = table.cell(r_idx, c_idx)
            cell.text = text.strip()
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Segoe UI'
                run.font.size = Pt(8.5 if not is_header else 9)
                if is_header:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    run.font.color.rgb = C_TEXT

            if is_header:
                set_cell_background(cell, '184C78')
            elif r_idx % 2 == 1:
                set_cell_background(cell, 'F8FAFC')
            else:
                set_cell_background(cell, 'FFFFFF')

            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(4)
    table_data = []
    in_table = False

in_code = False
code_lines = []

for line in lines:
    raw = line.rstrip('\r\n')
    
    # Code block toggle
    if raw.startswith('```'):
        if in_code:
            # flush code
            code_text = '\n'.join(code_lines)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 30, 30)
            code_lines = []
            in_code = False
        else:
            if in_table:
                flush_table()
            in_code = True
        continue

    if in_code:
        code_lines.append(raw)
        continue

    # Tables
    if raw.startswith('|') and raw.endswith('|'):
        parts = [c.strip() for c in raw.split('|')[1:-1]]
        if all(set(c) <= set('-: ') for c in parts):
            continue  # separator row
        in_table = True
        table_data.append(parts)
        continue
    else:
        if in_table:
            flush_table()

    if not raw.strip():
        continue

    # Headings
    if raw.startswith('# '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(raw[2:].strip())
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = C_PRIMARY
    elif raw.startswith('## '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(raw[3:].strip())
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = C_SECONDARY
    elif raw.startswith('### '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(raw[4:].strip())
        run.font.name = 'Segoe UI'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = C_PRIMARY
    elif raw.startswith('#### '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(raw[5:].strip())
        run.font.name = 'Segoe UI'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = C_TEXT
    elif raw.startswith('- ') or raw.startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(raw[2:].strip())
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9.5)
        run.font.color.rgb = C_TEXT
    elif raw.startswith('> '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.3)
        run = p.add_run(raw[2:].strip())
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9.5)
        run.font.italic = True
        run.font.color.rgb = C_TEXT
    elif raw == '---':
        continue
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(raw.strip())
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9.5)
        run.font.color.rgb = C_TEXT

if in_table:
    flush_table()

doc.save(docx_path)
print(f"Generated DOCX: {docx_path}")
