import os
from docx import Document
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EV = os.path.join(BASE, "evidence")


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def build_signoff_docx():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading(
        "BÁO CÁO BÀN GIAO CHÍNH THỨC (COMMERCIAL RELEASE SIGN-OFF): RATE-UI-05 → RATE-UI-07 (v2.0)",
        0,
    )
    doc.add_paragraph(
        "Ngày bàn giao: 17/10/2026 · Baseline kỹ thuật chốt: commit 5ec6653 · Commit đóng GAP DATA-01 (Seed danh mục chuẩn): commit 89f48e1 · Trạng thái nghiệm thu: 100/100 Commercial-Ready"
    )

    doc.add_heading("1. Kết quả thực hiện 3 công việc bàn giao cuối cùng của PO", 1)

    doc.add_heading("1.1 Công việc 1 — Seed Dữ liệu Khoản mục kinh tế ban đầu (Đóng GAP DATA-01)", 2)
    for item in [
        "Đã bổ sung seeder idempotent EconomicChargeCatalogSeeder (trong src/LCMS.Application/EconomicCharges/EconomicChargeCommands.cs), tự động chạy tại MigrateDatabaseAsync (src/LCMS.Api/Program.cs) cho tenant ops và demo, đồng thời tích hợp vào DemoDataSeeder.EnsureDefaultCatalogAsync.",
        "Đã cung cấp script SQL idempotent tại infra/sql/seed_economic_charge_types_ops_demo.sql (an toàn khi chạy nhiều lần trên PostgreSQL).",
        "Đã bổ sung kiểm thử tự động Data01_EconomicChargeCatalogSeeder_SeedsOpsAndDemoIdempotently trong tests/LCMS.Api.Tests/RateUi0507UatTests.cs (21/21 PASS).",
        "Đã thực thi và xác nhận 6/6 khoản mục chuẩn đều ở trạng thái IsActive = true trên cả local, cms-sg-01 (194.233.89.26) và cms-lss-01 (169.58.47.30):",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_table(
        doc,
        ["Mã khoản giá (Code)", "Tên hiển thị tiếng Việt (Name)", "Trạng thái (IsActive)", "Môi trường đã xác nhận"],
        [
            ["FREIGHT", "Cước vận tải chính (Base freight)", "true (Đang hoạt động)", "Local, cms-sg-01, cms-lss-01"],
            ["PICKUP", "Phí lấy hàng tận nơi", "true (Đang hoạt động)", "Local, cms-sg-01, cms-lss-01"],
            ["DELIVERY", "Phí giao hàng", "true (Đang hoạt động)", "Local, cms-sg-01, cms-lss-01"],
            ["REMOTE", "Phụ phí vùng xa / hải đảo", "true (Đang hoạt động)", "Local, cms-sg-01, cms-lss-01"],
            ["HANDLING", "Phí bốc xếp / nâng hạ", "true (Đang hoạt động)", "Local, cms-sg-01, cms-lss-01"],
            ["FUEL", "Phụ phí nhiên liệu", "true (Đang hoạt động)", "Local, cms-sg-01, cms-lss-01"],
        ],
        [3.2, 6.0, 3.3, 4.0],
    )

    doc.add_heading("1.2 Công việc 2 — Cập nhật tài liệu UAT Matrix chính thức (Baseline 5ec6653)", 2)
    for item in [
        "Đã cập nhật trực tiếp file ma trận chính thức LCMS_UAT_RATE_UI_05_07_DEV_UAT_Matrix_v1.0.xlsx và bản _RESULT.xlsx.",
        "Ghi nhận commit 5ec6653 là bản baseline kỹ thuật chốt cho gói RATE-UI-05..07 (và commit 89f48e1 đóng GAP DATA-01).",
        "Toàn bộ 20/20 kịch bản UAT (UAT-01 → UAT-20) đạt trạng thái PASS; toàn bộ 11/11 hạng mục DEV Checklist (UI-01 → DOC-01, bao gồm DATA-01) đạt trạng thái Hoàn thành (v2.0).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("1.3 Công việc 3 — Kiểm tra xác nhận triển khai trên VPS cms-sg-01 (194.233.89.26)", 2)
    add_table(
        doc,
        ["Hạng mục kiểm tra trên cms-sg-01", "Kết quả thực tế", "Bằng chứng"],
        [
            [
                "Trạng thái containers (docker ps)",
                "5/5 containers Up & Healthy",
                "cms-proxy (nginx:1.27-alpine), cms-api (infra-api), cms-web, cms-db (healthy), cms-redis đều Up.",
            ],
            [
                "Health & Readiness endpoints",
                "200 OK",
                'GET /health -> {"status":"ok","service":"cms-api"}; GET /ready -> {"status":"ready","database":"ok"}.',
            ],
            [
                "Database migration & Seed DATA-01 trên cms-db",
                "6/6 khoản mục chuẩn IsActive = true",
                "SELECT từ bảng economic_charge_types trên cms-db xác nhận đủ DELIVERY, FREIGHT, FUEL, HANDLING, PICKUP, REMOTE (is_active = t).",
            ],
            [
                "Container runtime logs (cms-api & cms-web)",
                "0 Error / 0 Fatal trên cms-api",
                "Log khởi động: 'No migrations were applied. The database is already up to date.' Outbox worker và API xử lý bình thường.",
            ],
        ],
        [4.5, 3.5, 8.5],
    )

    doc.add_heading("2. Tổng hợp 11/11 hạng mục DEV Checklist (Sau khi đóng GAP DATA-01)", 1)
    add_table(
        doc,
        ["ID", "Hạng mục", "Trạng thái", "Bằng chứng nghiệm thu v2.0"],
        [
            ["UI-01", "Ẩn loại giá ở cấp rule, basic và advanced", "Hoàn thành (v2.0)", "Đã bỏ ở Basic, Advanced và Mobile (rates.tsx)."],
            ["API-01", "Validate chiều giá rule và header", "Hoàn thành (v2.0)", "422 party_type_conflict trên mọi đường ghi; luôn lưu PartyTypeCode = null."],
            ["UI-02", "Disable/reset nấc trọng lượng", "Hoàn thành (v2.0)", "Khóa và xóa Từ–Đến kèm thông báo trên cả Web & Mobile."],
            ["API-02", "Validate bộ ba weight basis/min/max", "Hoàn thành (v2.0)", "WeightBand.DeclarationError trên Declare, AddPricingRule, Import, Publish."],
            ["UI-03", "Ẩn Nhóm cạnh tranh", "Hoàn thành (v2.0)", "Chọn Khoản giá từ master trên Basic, Advanced, Mobile; có nút Mở Khoản mục kinh tế."],
            ["DATA-01", "Xác minh danh mục khoản giá và mapping", "Hoàn thành (Đã đóng GAP)", "EconomicChargeCatalogSeeder + SQL script seed 6 khoản mục chuẩn (IsActive=true) cho ops & demo."],
            ["ENGINE-01", "Derive charge_code, chống cộng trùng", "Hoàn thành (v2.0)", "Suy ra từ master cho cả Basic & Advanced (xóa ReferenceCharge.Freight)."],
            ["ENGINE-02", "Không dùng tên/mã phân loại phụ phí cho version mới", "Hoàn thành (v2.0)", "Phiên bản explicit chỉ dùng lineRole; UAT-18 PASS."],
            ["MIG-01", "Dry-run và rollback, không sửa published", "Hoàn thành (v2.0)", "Không đổi schema; dữ liệu lịch sử bất biến; rollback bằng git revert."],
            ["QA-01", "20 UAT + regression BUY/SELL/Leg/advanced", "Hoàn thành (v2.0)", "21/21 RateUi0507UatTests, 415/415 dotnet test, 22/22 vitest, 45/45 mobile jest."],
            ["DOC-01", "PR, API contract, ảnh before/after, test logs", "Hoàn thành (v2.0)", "Baseline 5ec6653 + Seed 89f48e1, ảnh 1920px/768px/375px, UAT Matrix chính thức đã điền."],
        ],
        [1.8, 4.5, 3.0, 7.2],
    )

    doc.add_heading("3. Đề xuất ký biên bản bàn giao chính thức (Commercial Release Sign-off)", 1)
    doc.add_paragraph(
        "Toàn bộ yêu cầu kỹ thuật (Baseline commit 5ec6653) và dữ liệu danh mục Khoản mục kinh tế ban đầu (commit 89f48e1) "
        "đã hoàn tất 100% và hoạt động ổn định trên môi trường production cms-sg-01. Kính trình PO ký biên bản bàn giao chính thức (Commercial Release Sign-off)."
    )

    out_path = os.path.join(BASE, "LCMS_UAT_RATE_UI_05_07_Commercial_Release_SignOff_v2.0.docx")
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    print(build_signoff_docx())
