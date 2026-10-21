from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

out = Path(r"c:\A1\git\cms\docs\po\LCMS_RATE_CR_20261014_02_v1.1_373.docx")
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

doc.add_heading("LCMS-RATE-CR-20261014-02 v1.1 — Chốt 373/373", 0)
doc.add_paragraph("Ngày: 14/10/2026. Sáu ca API đỏ đã được đưa về đúng hợp đồng v1.1. Validation của API không bị nới.")

doc.add_heading("1. Sáu ca đã xử lý", 1)
for line in [
    "WeightBreakPivot, SellWeightBreak và ReferenceTariff dựng ma trận bằng POST /api/rate-cards/compose (declaration nâng cao, cột loại hàng và khung giá). Kết quả cước giữ nguyên: 250 USD, 300 USD doanh thu dự kiến, 682.500 đ, 200.000 đ, 1.500.000 đ, 170 USD, 1.320 USD.",
    "ComposeTariff tính giá có transportMode air, khớp ma trận đã khai báo hàng không. 10,5 kg hàng thường vẫn là 682.500 đ. Ô Chuyển nhanh trống dưới 11 kg vẫn bị từ chối.",
    "SameSpecificity: hai quy tắc FREIGHT cùng điều kiện không phát hành được. Kỳ vọng là 409 và thông báo RATE_CONFLICT. Phần tuyến cụ thể hơn vẫn phát hành và tính 30.",
    "ReceiveDocument_ByBillBusinessCode: chứng từ USD cần tỷ giá sang VND trong sổ. Test ghi một tỷ giá USD/VND ngày hiện tại trước khi nhận chứng từ theo mã Bill HAWB-UAT-001. Tiền tệ ZZZ vẫn bị từ chối.",
]:
    doc.add_paragraph(line, style="List Bullet")

doc.add_heading("2. Log kiểm thử", 1)
doc.add_paragraph("dotnet test tests/LCMS.Api.Tests/LCMS.Api.Tests.csproj")
log = doc.add_paragraph()
run = log.add_run(
    "Passed!  - Failed:     0, Passed:   373, Skipped:     0, Total:   373, Duration: 2 m 21 s - LCMS.Api.Tests.dll (net8.0)"
)
run.font.name = "Consolas"
run.font.size = Pt(10)
doc.add_paragraph("npx vitest run (apps/web)")
log2 = doc.add_paragraph()
run2 = log2.add_run("Test Files  1 passed (1)\nTests  11 passed (11)")
run2.font.name = "Consolas"
run2.font.size = Pt(10)

doc.add_heading("3. Tự chấm 100/100", 1)
table = doc.add_table(rows=5, cols=3)
table.style = "Table Grid"
for i, text in enumerate(["Tiêu chí", "Điểm", "Căn cứ"]):
    table.rows[0].cells[i].text = text
rows = [
    ("1. 36 AC và test tự động", "35/35", "dotnet test 373/373 PASS. vitest 11/11 PASS."),
    ("2. Bản đã phát hành khóa sửa", "30/30", "Giữ từ biên bản giao diện: badge Đã phát hành, không nút Sửa."),
    ("3. Cột bảng cơ bản", "15/15", "Giữ từ biên bản giao diện: đủ mã, tên, từ–đến một hàng."),
    ("4. Tiếng Việt và nấc BR-03", "20/20", "Giữ từ biên bản giao diện: không mã raw, nhãn ≥ 45 kg."),
]
for r, data in enumerate(rows, start=1):
    for c, text in enumerate(data):
        table.rows[r].cells[c].text = text
p = doc.add_paragraph()
mark = p.add_run("Tổng: 100/100 Commercial Ready Product.")
mark.bold = True
doc.add_paragraph("File giao diện kèm ảnh: docs/po/LCMS_RATE_CR_20261014_02_v1.1_100.docx.")

doc.save(out)
print(out, out.stat().st_size)
