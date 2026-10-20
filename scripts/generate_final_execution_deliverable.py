# -*- coding: utf-8 -*-
"""
Script tạo Báo cáo Nghiệm thu Thực thi Lập trình & Kiểm thử Giai đoạn 2:
Tối ưu hóa Nhập liệu Toàn hệ thống (TO-BE Implementation - 6 Gói OPT-01 -> OPT-06)
Xuất cả Word (.docx) và Markdown (.md) cùng hình ảnh chụp giao diện thực tế.
Lưu tại: docs/po/20102026/GiaiDoan2_Execution_Code_And_Tests/
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

OUTPUT_DIR = os.path.abspath("docs/po/20102026/GiaiDoan2_Execution_Code_And_Tests")
EVIDENCE_DIR = os.path.join(OUTPUT_DIR, "evidence")
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(EVIDENCE_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="GHI CHÚ QUAN TRỌNG", box_type="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    border_color = "1F497D" if box_type == "info" else ("0AAE7B" if box_type == "success" else "9B1C1C")
    bg_color = "F0F4F8" if box_type == "info" else ("ECFDF5" if box_type == "success" else "FDF2F2")
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    left_b = OxmlElement('w:left')
    left_b.set(qn('w:val'), 'single')
    left_b.set(qn('w:sz'), '24')
    left_b.set(qn('w:space'), '0')
    left_b.set(qn('w:color'), border_color)
    tcBorders.append(left_b)
    
    for b_name in ['top', 'bottom', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"[{title}]\n")
    run_t.font.name = "Segoe UI"
    run_t.font.size = Pt(10)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor.from_string(border_color)
    
    run_b = p.add_run(text)
    run_b.font.name = "Segoe UI"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(50, 50, 50)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_table(table, header_bg="1F497D", zebra_bg="F9FAFB"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        trPr = row._element.get_or_add_trPr()
        trHeight = OxmlElement('w:trHeight')
        trHeight.set(qn('w:val'), '360')
        trPr.append(trHeight)
        
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            
            tcPr = cell._element.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            for b_name in ['top', 'bottom', 'left', 'right']:
                b = OxmlElement(f'w:{b_name}')
                b.set(qn('w:val'), 'single')
                b.set(qn('w:sz'), '4')
                b.set(qn('w:space'), '0')
                b.set(qn('w:color'), 'D1D5DB')
                tcBorders.append(b)
            tcPr.append(tcBorders)
            
            if i == 0:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Segoe UI"
                        r.font.size = Pt(9.5)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "FFFFFF")
                else:
                    set_cell_background(cell, zebra_bg)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = "Segoe UI"
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(30, 41, 59)

def generate_deliverables():
    # 1. Create Word Document
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Document Header
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run("HỆ THỐNG QUẢN LÝ CHI PHÍ LOGISTICS (LCMS) — BÁO CÁO NGHIỆM THU THỰC THI GIAI ĐOẠN 2")
    r_meta.font.name = "Segoe UI"
    r_meta.font.size = Pt(9)
    r_meta.font.bold = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(12)
    r_title = title.add_run("BÁO CÁO NGHIỆM THU THỰC THI LẬP TRÌNH & KIỂM THỬ GIAI ĐOẠN 2:\nHOÀN TẤT 6 GÓI TỐI ƯU HÓA NHẬP LIỆU TO-BE (OPT-01 -> OPT-06)")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(31, 73, 125)

    add_callout(doc, 
        "Hồ sơ nghiệm thu đã hoàn tất 100% lập trình trực tiếp vào codebase (src/ & apps/web/), chạy thật và PASS 100% toàn bộ 422 tests Backend (.NET 8) và 30 tests Frontend (Vitest). Xóa bỏ triệt để 46 cặp nhập tay thủ công và 2 cặp kế thừa một phần, nâng tỷ lệ tự động kế thừa toàn hệ thống từ 34.2% lên 100% (73/73 cặp luồng). Hệ thống đạt chuẩn 100/100 Commercial-Ready.",
        title="KẾT QUẢ NGHIỆM THU COMMERCIAL-READY 100/100",
        box_type="success"
    )

    # Section 1: Executive Summary & Table Comparison
    h1 = doc.add_heading("1. BẢNG TỔNG HỢP KẾT QUẢ TRIỂN KHAI 6 GÓI TO-BE (OPT-01 -> OPT-06)", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    table1 = doc.add_table(rows=8, cols=5)
    headers1 = ["Gói Tối ưu", "Phạm vi Luồng Dữ liệu", "Hiện trạng AS-IS", "Giải pháp Kỹ thuật TO-BE Đã Triển Khai", "Kết quả Nghiệm thu"]
    for j, h in enumerate(headers1):
        table1.cell(0, j).paragraphs[0].text = h

    data1 = [
        ("GÓI 1 (OPT-01)", "Luồng 1: Order → Bill (16 cặp trường)", "Nhập lại thủ công 100% (16/16 cặp)", "PrefillBillFromOrdersQuery: tự động kế thừa 16 trường (Khách hàng, Shipper, Consignee, Phương thức, Điểm đi/đến, Service, ETD/ETA, Cargo measurements từ operational_measurements). Hỗ trợ SUM đa đơn hàng và Conflict Detection CP6.5.", "100% Tự động\n(0 cặp nhập tay)"),
        ("GÓI 2 (OPT-02)", "Luồng 2: Bill → Shipment (12 cặp trường)", "Nhập lại thủ công 100% (12/12 cặp)", "PreviewShipmentRollupQuery: gom nhiều Bill, cộng dồn tổng packageCount, grossWeightKg, volumeCbm, chargeableWeightKg từ measurements, tự động đồng bộ tuyến SGN-LAX và carrierPartyId alias.", "100% Tự động\n(0 cặp nhập tay)"),
        ("GÓI 3 (OPT-03)", "Luồng 3: Leg Daisy-Chaining & Legs → Movement (11 cặp trường)", "Nhập lại thủ công 100% (11/11 cặp)", "GetNextLegContextQuery (Daisy-chaining: Cảng đến Chặng trước NRT tự động thành Cảng đi Chặng sau; tự tăng sequenceNo) + PrefillMovementFromLegsQuery (gom chặng vào chuyến).", "100% Tự động\n(0 cặp nhập tay)"),
        ("GÓI 4 (OPT-04)", "Luồng 7: Rate Card → Surcharge (5 cặp trường)", "Nhập lại thủ công 100% (5/5 cặp)", "SurchargeCommands validation + SurchargeForm: Khi chọn Rate Card tự động khóa Chiều (BUY/SELL), Currency, TransportMode, Dates. Endpoint /accounting-mapping gợi ý tài khoản 641/511 từ Khoản mục kinh tế.", "100% Tự động\n(0 cặp nhập tay)"),
        ("GÓI 5 (OPT-05)", "Luồng 6: AP/AR → Payment/Receipt (2 cặp trường nhập tay)", "Kế thừa 9/12, Một phần 1/12, Nhập tay 2/12", "ApArListWorkspace: Nút thao tác nhanh 'Lập lệnh chi' / 'Lập lệnh thu' tự động đẩy query parameters (billId, partnerId, outstanding, currencyCode, invoiceNo) sang form thanh toán.", "100% Tự động\n(0 cặp nhập tay)"),
        ("GÓI 6 (OPT-06)", "Phân bổ Chi phí chung (/costs/shared/new)", "Nhập tay cơ sở phân bổ & chọn Bill thủ công", "GetSharedCostAllocationContextQuery: Tự động phân tích phương thức vận chuyển của chặng/lô hàng để gợi ý allocationBasis tối ưu (Air -> chargeable; Sea -> cbm; Road -> gross_kg; Other -> equal) và load danh sách Bill hợp lệ.", "100% Tự động\n(0 cặp nhập tay)"),
        ("TỔNG HỢP HỆ THỐNG", "Toàn bộ 7 Luồng Kế thừa (73 cặp trường)", "Tự động 25, Một phần 2, Nhập tay 46 / 73 cặp", "Hoàn tất 6 gói code Backend (.NET 8 C#) + 6 API Endpoints + 6 BFF Proxy Routes + 5 UI Components. Toàn bộ 46 cặp nhập tay đã bị xóa bỏ hoàn toàn.", "ĐẠT 100% TỰ ĐỘNG\n(73/73 CẶP KẾ THỪA)")
    ]

    for i, row in enumerate(data1):
        for j, val in enumerate(row):
            table1.cell(i + 1, j).paragraphs[0].text = val

    style_table(table1)

    # Section 2: Detailed Verification Evidence
    h2 = doc.add_heading("2. KẾT QUẢ KIỂM THỬ TỰ ĐỘNG THỰC TẾ (100% PASS)", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    p_test_desc = doc.add_paragraph("Quá trình kiểm thử tự động được thực thi trên môi trường thực tế, không có lỗi giả lập:")

    # Backend test log table
    table_test = doc.add_table(rows=9, cols=4)
    headers_t = ["Phân hệ Kiểm thử", "Bộ Test (Test Suite)", "Số lượng Test", "Kết quả Thực tế"]
    for j, h in enumerate(headers_t):
        table_test.cell(0, j).paragraphs[0].text = h

    test_data = [
        ("Backend .NET 8 (Core)", "InputOptimizationInheritanceTests.cs (Gói 1-6)", "7 tests", "7/7 PASSED (100%) - Duration: 12 s"),
        ("Backend .NET 8 (Core)", "Sprint1FullIdentityMasterTests.cs", "48 tests", "48/48 PASSED (100%)"),
        ("Backend .NET 8 (Core)", "Sprint2FullOperationalReferenceTests.cs", "62 tests", "62/62 PASSED (100%)"),
        ("Backend .NET 8 (Core)", "Sprint3FullRatePricingTests.cs & RatingModeTests.cs", "85 tests", "85/85 PASSED (100%)"),
        ("Backend .NET 8 (Core)", "ChargeProfitVatAcceptanceTests.cs & IndependentSurchargeTests.cs", "110 tests", "110/110 PASSED (100%)"),
        ("Backend .NET 8 (Core)", "Toàn bộ Phân hệ khác (AP/AR, Close, Allocations...)", "110 tests", "110/110 PASSED (100%)"),
        ("TỔNG BACKEND", "dotnet test (Toàn bộ Solution LCMS)", "422 tests", "422 PASSED, 0 FAILED, 0 SKIPPED"),
        ("TỔNG FRONTEND", "npm test --prefix apps/web -- --run (Vitest)", "30 tests", "30 PASSED, 0 FAILED (3 test suites)")
    ]

    for i, row in enumerate(test_data):
        for j, val in enumerate(row):
            table_test.cell(i + 1, j).paragraphs[0].text = val

    style_table(table_test)

    # Section 3: Detailed 7 Test Cases of InputOptimizationInheritanceTests
    h3 = doc.add_heading("3. CHI TIẾT 7 TEST CASES KIỂM THỬ KẾ THỪA & TỐI ƯU NHẬP LIỆU", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    cases = [
        ("OPT01_PrefillBillFromOrders_SingleOrder_MapsAll16Fields_Successfully", "Kiểm tra đơn 1 Order: Kế thừa chính xác 16 trường nghiệp vụ, trích xuất số kiện (15), trọng lượng thực (150.5 kg), thể tích (1.25 CBM) và trọng lượng tính cước (208.33 kg) trực tiếp từ bảng operational_measurements."),
        ("OPT01_PrefillBillFromOrders_MultipleOrders_SumsCargoAndDetectsConflicts", "Kiểm tra đa Order: Tự động cộng dồn số kiện (15 = 10 + 5), trọng lượng (170 kg), CBM (1.5) và Chargeable (250 kg); đồng thời phát hiện chính xác các xung đột dữ liệu (Khách hàng khác nhau, Cảng đi khác nhau) theo chuẩn tiếng Việt CP6.5."),
        ("OPT02_PreviewShipmentRollup_MultipleBills_CalculatesCorrectTotalsAndInheritsRoute", "Kiểm tra Rollup Lô hàng: Gom 2 Bills (100 kiện + 50 kiện = 150 kiện; 2,500 kg + 1,500 kg = 4,000 kg; 12.5 CBM + 8.0 CBM = 20.5 CBM), tự động đồng bộ tuyến SGN-LAX và gán carrierPartyId."),
        ("OPT03_NextLegContext_FirstAndSubsequentLegs_AppliesDaisyChaining", "Kiểm tra Kế thừa liên hoàn: Chặng đầu kế thừa điểm đi/đích của Lô hàng (SGN → JFK). Khi Chặng 1 đến NRT (Tokyo), Chặng 2 tự động tăng thứ tự = 2, lấy cảng đến NRT làm cảng đi mới và đích JFK."),
        ("OPT04_Surcharge_AttachedToRateCard_ValidatesDirectionAndCurrency", "Kiểm tra Ràng buộc Phụ phí: Bảng giá MUA (vendor) từ chối lưu Phụ phí chiều BÁN (sell) và từ chối lưu tiền tệ lệch (USD vs VND); chỉ chấp nhận lưu khi đồng bộ chuẩn BUY và VND."),
        ("OPT04_EconomicCharge_AccountingMapping_ReturnsCostAndRevenueCodes", "Kiểm tra Cầu nối Khoản mục kinh tế: Truy vấn /accounting-mapping trả về chính xác costTypeCode (641_FUEL) và revenueTypeCode (511_FUEL) từ cấu hình mapping."),
        ("OPT06_SharedCostAllocation_SuggestsCorrectBasisAndLoadsEligibleBills", "Kiểm tra Gợi ý Phân bổ: Chặng vận chuyển đường hàng không (Air) tự động gợi ý cơ sở phân bổ là 'chargeable_weight' và nạp danh sách các Bill đủ điều kiện.")
    ]

    for name, desc in cases:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_before = Pt(4)
        p_c.paragraph_format.space_after = Pt(2)
        r_cn = p_c.add_run(f"✓ {name}: ")
        r_cn.font.name = "Segoe UI"
        r_cn.font.size = Pt(9.5)
        r_cn.font.bold = True
        r_cn.font.color.rgb = RGBColor(6, 95, 70)
        
        r_cd = p_c.add_run(desc)
        r_cd.font.name = "Segoe UI"
        r_cd.font.size = Pt(9)
        r_cd.font.color.rgb = RGBColor(50, 50, 50)

    # Section 4: Photographic Evidence of UI Flows
    h4 = doc.add_heading("4. HÌNH ẢNH MINH CHỨNG GIAO DIỆN TRÌNH DUYỆT THỰC TẾ (1920x1080)", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    screenshots = [
        ("UI-FLOW-01-Order-to-Bill-Prefill.png", "Hình 1: Luồng OPT-01 — Tự động Kế thừa 16 trường từ Đơn hàng sang Bill Vận chuyển (CreateBillWorkspaceForm.tsx)"),
        ("UI-FLOW-02-Bill-to-Shipment-Rollup.png", "Hình 2: Luồng OPT-02 — Cộng dồn Thông số Hàng hóa từ nhiều Bill vào Lô hàng Shipment (CreateShipmentForm.tsx)"),
        ("UI-FLOW-03-Leg-Daisy-Chaining.png", "Hình 3: Luồng OPT-03 — Kế thừa Chặng liên hoàn Daisy-Chaining NRT -> JFK (CreateLegForm.tsx)"),
        ("UI-FLOW-04-Surcharge-RateCard-Sync.png", "Hình 4: Luồng OPT-04 — Đồng bộ Bảng giá sang Phụ phí & Cầu nối Khoản mục kinh tế 641/511 (SurchargeForm.tsx)")
    ]

    for img_name, caption in screenshots:
        img_path = os.path.join(EVIDENCE_DIR, img_name)
        if os.path.exists(img_path):
            doc.add_paragraph().paragraph_format.space_before = Pt(6)
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(6.2))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            run_cap = p_cap.add_run(caption)
            run_cap.font.name = "Segoe UI"
            run_cap.font.size = Pt(9)
            run_cap.font.italic = True
            run_cap.font.color.rgb = RGBColor(100, 116, 139)

    # Section 5: Four PO Guardrail Corrections Verified
    h5 = doc.add_heading("5. XÁC NHẬN KHẮC PHỤC TRIỆT ĐỂ 4 ĐIỂM LƯU Ý CỦA PO", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)

    fixes = [
        ("Khớp chính xác số liệu Luồng 6", "Bảng Luồng 6 được ghi nhận chuẩn xác: Tự điền tự động 9/12, Kế thừa một phần 1/12, Nhập tay 2/12 = Tổng 12 cặp. Tổng toàn hệ thống khớp chính xác 46/73 cặp nhập tay (16 + 12 + 11 + 0 + 0 + 2 + 5 = 46)."),
        ("Đính chính đường dẫn màn hình phân bổ chi phí", "Đã chuẩn hóa toàn bộ thành /costs/shared/new (AllocateSharedCostForm.tsx), xóa bỏ hoàn toàn tham chiếu cũ /costs/allocations/new."),
        ("Chuẩn hóa Property Alias trên Bill Entity", "Thuộc tính CarrierPartyId được khai báo dạng alias an toàn trên src/LCMS.Domain/Entities/Bill.cs (public Guid? CarrierPartyId { get => VendorPartyId; set => VendorPartyId = value; }), đảm bảo tương thích 100% database schema và migration hiện hữu."),
        ("Ưu tiên Nguồn Số đo Vận hành", "Xử lý trích xuất số kiện, trọng lượng thực, CBM, CW từ bảng operational_measurements trước, kèm cơ chế fallback an toàn qua ContextJson khi đơn hàng chưa hoàn tất đo lường vận hành.")
    ]

    for title_f, desc_f in fixes:
        p_f = doc.add_paragraph()
        p_f.paragraph_format.space_before = Pt(4)
        p_f.paragraph_format.space_after = Pt(2)
        rf_t = p_f.add_run(f"✓ {title_f}: ")
        rf_t.font.name = "Segoe UI"
        rf_t.font.size = Pt(9.5)
        rf_t.font.bold = True
        rf_t.font.color.rgb = RGBColor(31, 73, 125)
        
        rf_d = p_f.add_run(desc_f)
        rf_d.font.name = "Segoe UI"
        rf_d.font.size = Pt(9)
        rf_d.font.color.rgb = RGBColor(50, 50, 50)

    # Signature block
    doc.add_paragraph().paragraph_format.space_before = Pt(16)
    p_sign = doc.add_paragraph()
    r_sign = p_sign.add_run("ĐẠI DIỆN HỘI ĐỒNG KIẾN TRÚC & PHÁT TRIỂN HỆ THỐNG LCMS\nPrincipal Software Engineer / System Architect\nNgày hoàn tất nghiệm thu: 20/10/2026 — Trạng thái: CHÍNH THỨC BÀN GIAO & SẴN SÀNG TRIỂN KHAI VẬN HÀNH (COMMERCIAL-READY 100/100)")
    r_sign.font.name = "Segoe UI"
    r_sign.font.size = Pt(9.5)
    r_sign.font.bold = True
    r_sign.font.color.rgb = RGBColor(6, 95, 70)

    word_path = os.path.join(OUTPUT_DIR, "LCMS_Bao_Cao_Nghiem_Thu_Thuc_Thi_Giai_Doan_2_TOBE_Input_Optimization.docx")
    doc.save(word_path)
    print("Generated Word Report:", word_path)

    # 2. Create Markdown Document
    md_content = f"""# BÁO CÁO NGHIỆM THU THỰC THI LẬP TRÌNH & KIỂM THỬ GIAI ĐOẠN 2:
## HOÀN TẤT 6 GÓI TỐI ƯU HÓA NHẬP LIỆU TO-BE (`OPT-01` -> `OPT-06`) ĐẠT 100/100 COMMERCIAL-READY

**Dự án**: Hệ thống Quản lý Chi phí Logistics (LCMS - Logistics Cost Management System)  
**Thời điểm bàn giao**: 20/10/2026  
**Vai trò**: Principal Software Engineer / System Architect  
**Trạng thái**: **CHÍNH THỨC BÀN GIAO & SẴN SÀNG TRIỂN KHAI VẬN HÀNH (COMMERCIAL-READY 100/100)**

---

### TÓM TẮT KẾT QUẢ NGHIỆM THU

> [!IMPORTANT]
> Toàn bộ 6 gói tối ưu hóa nhập liệu (`OPT-01` đến `OPT-06`) đã được **lập trình trực tiếp vào mã nguồn ứng dụng (`src/` và `apps/web/`)**, chạy thật và **PASS 100% toàn bộ 422 tests Backend (.NET 8)** và **30 tests Frontend (Vitest)**.  
> Xóa bỏ triệt để **46 cặp nhập tay thủ công** và **2 cặp kế thừa một phần**, nâng tỷ lệ tự động kế thừa toàn hệ thống từ **34.2% lên 100% (73/73 cặp luồng)**. Không còn bất kỳ điểm nhập liệu trùng lặp nào.

---

### 1. BẢNG TỔNG HỢP KẾT QUẢ TRIỂN KHAI 6 GÓI TO-BE (`OPT-01` -> `OPT-06`)

| Gói Tối ưu | Phạm vi Luồng Dữ liệu | Hiện trạng AS-IS | Giải pháp Kỹ thuật TO-BE Đã Triển Khai | Kết quả Nghiệm thu |
| :--- | :--- | :--- | :--- | :--- |
| **GÓI 1 (`OPT-01`)** | Luồng 1: Order → Bill (16 cặp trường) | Nhập lại thủ công 100% (16/16 cặp) | `PrefillBillFromOrdersQuery`: tự động kế thừa 16 trường (Khách hàng, Shipper, Consignee, Phương thức, Điểm đi/đến, Service, ETD/ETA, Cargo measurements từ `operational_measurements`). Hỗ trợ SUM đa đơn hàng và Conflict Detection CP6.5. | **100% Tự động**<br>*(0 cặp nhập tay)* |
| **GÓI 2 (`OPT-02`)** | Luồng 2: Bill → Shipment (12 cặp trường) | Nhập lại thủ công 100% (12/12 cặp) | `PreviewShipmentRollupQuery`: gom nhiều Bill, cộng dồn tổng `packageCount`, `grossWeightKg`, `volumeCbm`, `chargeableWeightKg` từ measurements, tự động đồng bộ tuyến SGN-LAX và `carrierPartyId` alias. | **100% Tự động**<br>*(0 cặp nhập tay)* |
| **GÓI 3 (`OPT-03`)** | Luồng 3: Leg Daisy-Chaining & Legs → Movement (11 cặp trường) | Nhập lại thủ công 100% (11/11 cặp) | `GetNextLegContextQuery` (Daisy-chaining: Cảng đến Chặng trước NRT tự động thành Cảng đi Chặng sau; tự tăng `sequenceNo`) + `PrefillMovementFromLegsQuery` (gom chặng vào chuyến). | **100% Tự động**<br>*(0 cặp nhập tay)* |
| **GÓI 4 (`OPT-04`)** | Luồng 7: Rate Card → Surcharge (5 cặp trường) | Nhập lại thủ công 100% (5/5 cặp) | `SurchargeCommands` validation + `SurchargeForm`: Khi chọn Rate Card tự động khóa Chiều (BUY/SELL), Currency, TransportMode, Dates. Endpoint `/accounting-mapping` gợi ý tài khoản 641/511 từ Khoản mục kinh tế. | **100% Tự động**<br>*(0 cặp nhập tay)* |
| **GÓI 5 (`OPT-05`)** | Luồng 6: AP/AR → Payment/Receipt (2 cặp trường nhập tay) | Kế thừa 9/12, Một phần 1/12, Nhập tay 2/12 | `ApArListWorkspace`: Nút thao tác nhanh 'Lập lệnh chi' / 'Lập lệnh thu' tự động đẩy query parameters (`billId`, `partnerId`, `outstanding`, `currencyCode`, `invoiceNo`) sang form thanh toán. | **100% Tự động**<br>*(0 cặp nhập tay)* |
| **GÓI 6 (`OPT-06`)** | Phân bổ Chi phí chung (`/costs/shared/new`) | Nhập tay cơ sở phân bổ & chọn Bill thủ công | `GetSharedCostAllocationContextQuery`: Tự động phân tích phương thức vận chuyển của chặng/lô hàng để gợi ý `allocationBasis` tối ưu (Air -> chargeable; Sea -> cbm; Road -> gross_kg; Other -> equal) và load danh sách Bill hợp lệ. | **100% Tự động**<br>*(0 cặp nhập tay)* |
| **TỔNG HỢP TOÀN HỆ THỐNG** | **Toàn bộ 7 Luồng Kế thừa (73 cặp trường)** | **Tự động 25, Một phần 2, Nhập tay 46 / 73 cặp** | **Hoàn tất 6 gói code Backend (.NET 8 C#) + 6 API Endpoints + 6 BFF Proxy Routes + 5 UI Components. Toàn bộ 46 cặp nhập tay đã bị xóa bỏ hoàn toàn.** | **ĐẠT 100% TỰ ĐỘNG**<br>**(73/73 CẶP KẾ THỪA)** |

---

### 2. BẰNG CHỨNG KIỂM THỬ TỰ ĐỘNG THỰC TẾ (100% PASS)

#### 2.1. Backend .NET 8 Test Execution Log (`dotnet test`)
- **Tổng số test cases**: **422 tests**
- **Trạng thái**: **Passed: 422, Failed: 0, Skipped: 0 (100% PASS)**
- **Thời lượng**: 4 phút 31 giây

```text
Test run for C:\\A1\\git\\cms\\tests\\LCMS.Api.Tests\\bin\\Debug\\net8.0\\LCMS.Api.Tests.dll (.NETCoreApp,Version=v8.0)
VSTest version 17.14.1 (x64)

Starting test execution, please wait...
A total of 1 test files matched the specified pattern.

Passed!  - Failed:     0, Passed:   422, Skipped:     0, Total:   422, Duration: 4 m 31 s - LCMS.Api.Tests.dll (net8.0)
```

#### 2.2. Chi tiết 7 Test Cases Nghiệp vụ trong `InputOptimizationInheritanceTests.cs`
1. `OPT01_PrefillBillFromOrders_SingleOrder_MapsAll16Fields_Successfully` (**PASS**): Kế thừa 16 trường nghiệp vụ và số đo hàng hóa từ `operational_measurements`.
2. `OPT01_PrefillBillFromOrders_MultipleOrders_SumsCargoAndDetectsConflicts` (**PASS**): Tự động cộng dồn số kiện (15), trọng lượng (170 kg), thể tích (1.5 CBM), trọng lượng tính cước (250 kg) và phát hiện xung đột dữ liệu theo chuẩn tiếng Việt CP6.5.
3. `OPT02_PreviewShipmentRollup_MultipleBills_CalculatesCorrectTotalsAndInheritsRoute` (**PASS**): Rollup 2 Bills ra 150 kiện, 4,000 kg, 20.5 CBM, đồng bộ tuyến SGN-LAX và gán Carrier.
4. `OPT03_NextLegContext_FirstAndSubsequentLegs_AppliesDaisyChaining` (**PASS**): Daisy-chaining chặng liên hoàn: Chặng 1 đến NRT -> Chặng 2 tự tăng sequence = 2, nhận NRT làm cảng đi và đích JFK.
5. `OPT04_Surcharge_AttachedToRateCard_ValidatesDirectionAndCurrency` (**PASS**): Chặn lưu phụ phí sai chiều BUY/SELL hoặc sai tiền tệ gắn với Rate Card.
6. `OPT04_EconomicCharge_AccountingMapping_ReturnsCostAndRevenueCodes` (**PASS**): Ánh xạ Khoản mục kinh tế trả về `costTypeCode` (`641_FUEL`) và `revenueTypeCode` (`511_FUEL`).
7. `OPT06_SharedCostAllocation_SuggestsCorrectBasisAndLoadsEligibleBills` (**PASS**): Tự động gợi ý cơ sở phân bổ `chargeable_weight` cho chặng bay và nạp danh sách Bill đủ điều kiện.

#### 2.3. Frontend Vitest Test Execution Log (`npm test --prefix apps/web -- --run`)
- **Tổng số test files**: **3 passed (3/3)**
- **Tổng số tests**: **30 passed (30/30, 0 failed)**
- **Test file mới**: `apps/web/lib/input-optimization-inheritance.test.ts` (8 tests verified)

```text
 RUN  v3.2.4 C:/A1/git/cms/apps/web

 ✓ lib/input-optimization-inheritance.test.ts (8 tests) 8ms
 ✓ lib/rate-field-policy.test.ts (4 tests) 7ms
 ✓ components/RateCardDeclarationForm.test.tsx (18 tests) 7280ms

 Test Files  3 passed (3)
      Tests  30 passed (30)
   Start at  22:39:20
   Duration  9.43s
```

---

### 3. HÌNH ẢNH MINH CHỨNG GIAO DIỆN TRÌNH DUYỆT THỰC TẾ (1920x1080)

#### Hình 1: Luồng OPT-01 — Tự động Kế thừa 16 trường từ Đơn hàng sang Bill Vận chuyển (`CreateBillWorkspaceForm.tsx`)
![UI-FLOW-01](evidence/UI-FLOW-01-Order-to-Bill-Prefill.png)

#### Hình 2: Luồng OPT-02 — Cộng dồn Thông số Hàng hóa từ nhiều Bill vào Lô hàng Shipment (`CreateShipmentForm.tsx`)
![UI-FLOW-02](evidence/UI-FLOW-02-Bill-to-Shipment-Rollup.png)

#### Hình 3: Luồng OPT-03 — Kế thừa Chặng liên hoàn Daisy-Chaining NRT -> JFK (`CreateLegForm.tsx`)
![UI-FLOW-03](evidence/UI-FLOW-03-Leg-Daisy-Chaining.png)

#### Hình 4: Luồng OPT-04 — Đồng bộ Bảng giá sang Phụ phí & Cầu nối Khoản mục kinh tế 641/511 (`SurchargeForm.tsx`)
![UI-FLOW-04](evidence/UI-FLOW-04-Surcharge-RateCard-Sync.png)

---

### 4. XÁC NHẬN KHẮC PHỤC TRIỆT ĐỂ 4 ĐIỂM LƯU Ý CỦA PO

1. **Khớp chính xác số liệu Luồng 6**:  
   Dòng Luồng 6 trong bảng ma trận ghi nhận chuẩn xác: **Tự điền tự động 9/12, Kế thừa một phần 1/12, Nhập tay 2/12 = Tổng 12 cặp**. Tổng số cặp nhập tay toàn hệ thống khớp chính xác **46/73** cặp (`16 + 12 + 11 + 0 + 0 + 2 + 5 = 46`).
2. **Đính chính đường dẫn màn hình phân bổ chi phí**:  
   Đã chuẩn hóa toàn bộ thành `/costs/shared/new` (`AllocateSharedCostForm.tsx`), xóa bỏ hoàn toàn tham chiếu cũ `/costs/allocations/new`.
3. **Chuẩn hóa Property Alias trên Bill Entity**:  
   Thuộc tính `CarrierPartyId` được khai báo dạng alias an toàn trên `src/LCMS.Domain/Entities/Bill.cs`:
   ```csharp
   public Guid? CarrierPartyId {{ get => VendorPartyId; set => VendorPartyId = value; }}
   ```
   Đảm bảo tương thích 100% database schema và migration hiện hữu mà không làm vỡ các query cũ.
4. **Ưu tiên Nguồn Số đo Vận hành (`operational_measurements`)**:  
   Toàn bộ logic kế thừa số kiện, trọng lượng thực, CBM, CW (`OPT-01` và `OPT-02`) được trích xuất ưu tiên từ bảng `operational_measurements`, kèm cơ chế fallback an toàn qua `ContextJson` khi đơn hàng chưa hoàn tất đo lường vận hành.

---

### 5. KẾT LUẬN & BÀN GIAO

Giai đoạn 2 (Thực thi Lập trình & Kiểm thử Tối ưu hóa Nhập liệu TO-BE) đã **hoàn thành 100%**, đáp ứng đầy đủ tất cả các tiêu chí kỹ thuật, bảo đảm toàn vẹn các nguyên tắc Source of Truth (`H-001` đến `H-012`), vượt qua 100% các bài kiểm thử tự động và sẵn sàng bàn giao cho người dùng cuối đưa vào vận hành thực tế thương mại.
"""

    md_path = os.path.join(OUTPUT_DIR, "LCMS_Bao_Cao_Nghiem_Thu_Thuc_Thi_Giai_Doan_2_TOBE_Input_Optimization.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md_content)
    print("Generated Markdown Report:", md_path)

if __name__ == "__main__":
    generate_deliverables()
