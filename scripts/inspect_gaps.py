import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('docs/po/21102026-LCMS_BA_DataEntry_TO_BE_v0.9_Package/LCMS_TO_BE_Field_Optimization_Inheritance_Matrix_v0.9.xlsx', data_only=True)
ws5 = wb['05 GAP và rủi ro']

for r in range(1, ws5.max_row + 1):
    vals = [str(ws5.cell(r, c).value or '').strip().replace('\n', ' ') for c in range(1, ws5.max_column + 1)]
    print(f"[{vals[0]}] {vals[1]} | Phát hiện: {vals[2]} | DEV cần làm: {vals[3]} | Mức độ: {vals[4]} | Tình trạng: {vals[5]}")
