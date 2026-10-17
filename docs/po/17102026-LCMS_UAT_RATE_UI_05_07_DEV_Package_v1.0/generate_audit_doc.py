import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_borders(cell, top="CCCCCC", bottom="CCCCCC", left=None, right=None, left_sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, color in borders.items():
        if color:
            sz = left_sz if edge == 'left' and left_sz > 4 else 4
            elm = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(elm)
        else:
            elm = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none" w:sz="0" w:space="0" w:color="auto"/>')
            tcBorders.append(elm)
    tcPr.append(tcBorders)

def add_callout(doc, text, title=None, callout_type="warning"):
    colors = {
        "warning": {"bg": "FFF8F0", "border": "DD6B20", "title": "C05621"},
        "danger": {"bg": "FFF5F5", "border": "E53E3E", "title": "C53030"},
        "info": {"bg": "F0F7FF", "border": "2B6CB0", "title": "2B6CB0"},
        "success": {"bg": "F0FFF4", "border": "38A169", "title": "276749"},
        "prompt": {"bg": "F8FAFC", "border": "4A5568", "title": "2D3748"}
    }
    c = colors.get(callout_type, colors["info"])
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, c["bg"])
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_borders(cell, top=None, bottom=None, right=None, left=c["border"], left_sz=24)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(f"[{title.upper()}]\n")
        run_title.bold = True
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(10.5)
        run_title.font.color.rgb = RGBColor.from_string(c["title"])
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor.from_string("2D3748")
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def format_table(tbl, col_widths, headers, rows):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Format header row
    hdr_row = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        set_cell_borders(cell, top="1B365D", bottom="1B365D", left="2B4C7E", right="2B4C7E")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Format body rows
    for r_idx, r_data in enumerate(rows):
        row = tbl.add_row()
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            set_cell_borders(cell, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor.from_string("2D3748")
                
    # Set column widths
    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    run.font.name = "Calibri"
    if level == 1:
        run.font.size = Pt(14)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("1B365D")
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(12)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("2B6CB0")
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(10.5)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("2D3748")
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor.from_string("1A202C")
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor.from_string("2D3748")
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor.from_string("1A202C")
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor.from_string("2D3748")
    return p

def main():
    doc = docx.Document()
    
    # Set page margins: 0.75 in (approx 1.9 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.different_first_page_header_footer = False
        
    # Title & Subtitle
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("LCMS — BÁO CÁO PHẢN BIỆN KẾT QUẢ UAT & CHỈ THỊ TÁI THẨM ĐỊNH KHẮT KHE")
    t_run.font.name = "Calibri"
    t_run.font.size = Pt(18)
    t_run.bold = True
    t_run.font.color.rgb = RGBColor.from_string("1B365D")
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(8)
    sub_run = sub_p.add_run("Thẩm tra kết quả thực hiện RATE-UI-05 → RATE-UI-07 (Commit bfd3d9e) | Kèm Prompt Tái chấm điểm 100 điểm cho Dev AI")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(10.5)
    sub_run.italic = True
    sub_run.font.color.rgb = RGBColor.from_string("4A5568")

    # Meta Table
    tbl_meta = doc.add_table(rows=1, cols=4)
    format_table(tbl_meta, [1.5, 2.0, 1.5, 2.0], 
                 ["Thông tin", "Giá trị", "Thông tin", "Giá trị"],
                 [
                     ["Ngày thẩm định", "17/10/2026", "Phiên bản tài liệu", "v1.0 (Audit & Directive)"],
                     ["Hạng mục thẩm tra", "RATE-UI-05, 06, 07", "Commit đối chiếu", "bfd3d9e (nhánh main)"],
                     ["Căn cứ nghiệp vụ", "LCMS_BA_UAT_RATE_UI_05_07_DEV_Spec_v1.0", "Trạng thái", "YÊU CẦU TÁI CHẤM ĐIỂM (CONDITIONAL)"]
                 ])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # PHẦN 1
    # -------------------------------------------------------------
    add_heading_styled(doc, "PHẦN 1: TỔNG QUAN YÊU CẦU PO VÀ BÁO CÁO CỦA DEV AI", level=1)
    
    add_p(doc, "Gói yêu cầu điều chỉnh sau UAT (RATE-UI-05..07) được Product Owner (PO) ban hành ngày 17/10/2026 với 3 Change Requests cốt lõi nhằm giải quyết triệt để các xung đột dữ liệu tài chính, đảm bảo tính chặt chẽ trong khai báo và tính toán giá của phân hệ LCMS:")
    
    add_bullet(doc, "Giữ duy nhất trường Loại giá (Giá mua/BUY=vendor; Giá bán/SELL=customer) tại Thông tin chung (rate_cards.party_type). Loại bỏ hoàn toàn dropdown Loại giá ở cấp dòng quy tắc (basic và advanced). API từ chối nếu client gửi giá trị trái header; Rating Engine không bị rule đảo chiều tài chính.", "1. RATE-UI-05 (Loại giá chỉ ở Thông tin chung): ")
    add_bullet(doc, "Đổi nhãn thành 'Điều kiện trọng lượng' gồm 3 chế độ: Không lọc nấc, Trọng lượng thực (gross), Trọng lượng tính cước (chargeable). Khi 'Không lọc nấc': khóa Từ/Đến, reset null cả 3 trường weightBasis/weightMin/weightMax. Khi có lọc: bắt buộc nhập đủ Từ và Đến (Từ ≥ 0, Đến > Từ). Basic áp dụng nửa khoảng [Từ, Đến); Advanced giữ upper bound inclusive.", "2. RATE-UI-06 (Điều kiện trọng lượng và nấc Từ – Đến): ")
    add_bullet(doc, "Ẩn ô nhập kỹ thuật 'Nhóm cạnh tranh' (charge_code) khỏi UI. Nhóm cạnh tranh phải được hệ thống xác định tự động từ định danh Khoản giá chuẩn (economic_charge_types / danh mục Khoản mục kinh tế), không tự gán FREIGHT. Thiếu master thì cho lưu nháp nhưng chặn phát hành. Giữ nguyên 5 quy tắc REMOTE/DELIVERY lịch sử đã phát hành và cảnh báo nguy cơ.", "3. RATE-UI-07 (Ẩn Nhóm cạnh tranh, hệ thống xác định): ")

    add_p(doc, "Báo cáo của Dev AI (tại file LCMS_UAT_RATE_UI_05_07_DEV_Result_v1.0.docx và commit bfd3d9e) tuyên bố đã hoàn thành 100%:", bold_prefix="Tuyên bố từ phía Dev AI: ")
    add_bullet(doc, "Toàn bộ 20/20 kịch bản UAT đều đánh dấu PASS kèm theo test tự động và log kiểm thử.")
    add_bullet(doc, "Bộ kiểm thử backend đạt 414/414 PASS (0 skipped); frontend vitest đạt 20/20 PASS; tsc --noEmit không lỗi.")
    add_bullet(doc, "Không sửa đổi DB schema, không chạy EF migration, không drop cột, không sửa đè lịch sử rating và published versions.")
    add_bullet(doc, "Chỉ thừa nhận 1 GAP duy nhất là DATA-01: Tenant demo chưa có danh mục Khoản mục kinh tế, bảng giá mới chỉ lưu được nháp, không phát hành được cho tới khi khai báo danh mục.")

    # -------------------------------------------------------------
    # PHẦN 2
    # -------------------------------------------------------------
    add_heading_styled(doc, "PHẦN 2: BÁO CÁO THẨM TRA ĐỐI CHIẾU & VẠCH TÌM ĐIỂM MÙ (AUDIT FINDINGS)", level=1)
    
    add_p(doc, "Qua quá trình kiểm tra chuyên sâu mã nguồn (code review), đối chiếu hợp đồng API, schema dữ liệu và phân tích logic kiểm thử, Ban thẩm định ghi nhận các nỗ lực đáng kể nhưng đồng thời PHÁT HIỆN 7 LỖ HỔNG / ĐIỂM MÙ TRỌNG YẾU bị Dev AI che lấp hoặc bỏ sót:")

    add_callout(doc, 
                "Dev AI đã tự chấm điểm 'PASS 100%' dựa trên các bài test được viết riêng (RateUi0507UatTests.cs). Tuy nhiên, các bài test này chỉ đi theo 'happy path' của riêng màn hình tổng hợp web, hoàn toàn bỏ qua các API endpoint nhánh, màn hình Mobile, tính năng Import và các kịch bản biên thực tế. Nếu đưa lên Production, rủi ro lỗi dữ liệu tài chính và vỡ giao diện là rất cao.",
                title="CẢNH BÁO TỪ BAN THẨM ĐỊNH", 
                callout_type="danger")

    add_heading_styled(doc, "2.1. Đánh giá khách quan các điểm tích cực", level=2)
    add_bullet(doc, "Thiết kế kế thừa BUY/SELL từ header rất rõ ràng, Rating Engine tuân thủ nguyên tắc financial_nature (BUY luôn ra cost, SELL luôn ra revenue), không bị rule đảo ngược tài chính.", "Tuân thủ triết lý tài chính: ")
    add_bullet(doc, "Đã bảo toàn nguyên vẹn tính bất biến (immutability) của dữ liệu đã phát hành và lịch sử rating_details.rule_code. Không phát sinh EF Core migration thừa.", "Bảo toàn dữ liệu lịch sử: ")
    add_bullet(doc, "Tạo mới query GET /api/rate-versions/{id}/rule-warnings giúp người dùng nhìn thấy cảnh báo về 5 quy tắc REMOTE/DELIVERY của NSE và các nấc trọng lượng lệch chuẩn cũ.", "Tính năng cảnh báo hữu ích: ")

    add_heading_styled(doc, "2.2. Vạch lá tìm sâu — 7 Lỗ hổng và Điểm mù kỹ thuật nghiêm trọng", level=2)

    # Lỗ hổng 1
    add_heading_styled(doc, "Lỗ hổng 1: Lỗ hổng Validation trên API thêm quy tắc đơn lẻ (AddPricingRuleCommand)", level=3)
    add_bullet(doc, "Tại src/LCMS.Application/PricingRules/Commands/PricingRuleCommands.cs (AddPricingRuleCommandHandler), hệ thống HOÀN TOÀN KHÔNG gọi WeightBand.DeclarationError(request.WeightBasis, request.WeightMin, request.WeightMax)! Nếu client hoặc hệ thống ngoài gọi POST /api/rate-versions/{id}/rules với weightBasis=null nhưng min=0, max=5 (vi phạm UAT-07) hoặc min=10, max=5 (vi phạm UAT-10), API vẫn chấp nhận lưu vào database mà không trả lỗi 422!", bold_prefix="Thiếu kiểm tra trọng lượng: ")
    add_bullet(doc, "Khi client gửi partyTypeCode = 'vendor' trên bảng giá 'vendor', AddPricingRuleCommand lưu trực tiếp 'vendor' vào DB thay vì lưu null như quy định cho bản nháp mới, gây bất nhất với SaveRateCardDeclarationCommand.", bold_prefix="Bất nhất lưu trữ party_type_code: ")
    add_bullet(doc, "AddPricingRuleCommand không có trường EconomicChargeTypeId trong request DTO, chỉ nhận chargeCode tự do dạng chuỗi.", bold_prefix="Khuyết thiếu liên kết master: ")

    # Lỗ hổng 2
    add_heading_styled(doc, "Lỗ hổng 2: Bỏ quên hoàn toàn ứng dụng Di động (Mobile App)", level=3)
    add_bullet(doc, "Tài liệu BA Spec mục 4 ghi rõ: 'Kiểm tra cùng hành vi trên màn hình tạo, sửa, copy, import, advanced và mobile/table responsive.'", bold_prefix="Vi phạm phạm vi đặc tả PO: ")
    add_bullet(doc, "Trong mã nguồn apps/mobile/app/modules/rates.tsx (hàm handleAddPricingRule dòng 338), app gọi POST /api/rate-versions/{id}/rules nhưng KHÔNG hề có trường chọn Khoản giá (economicChargeTypeId), KHÔNG có chargeCode, KHÔNG có điều kiện trọng lượng. Khi người dùng thêm quy tắc từ Mobile, bảng giá sẽ chứa quy tắc thiếu Khoản giá và sau đó BỊ KẸT CHẶN PHÁT HÀNH (409) trên Web mà không có cách nào xử lý trên app di động!", bold_prefix="Phát sinh 'Deadlock' nghiệp vụ: ")

    # Lỗ hổng 3
    add_heading_styled(doc, "Lỗ hổng 3: Bảng giá Nâng cao (Ma trận cước) vẫn âm thầm hardcode FREIGHT và bypass master", level=3)
    add_bullet(doc, "Dev AI khẳng định: 'Chủ đích của CR-07: không tự gán FREIGHT. Tenant chưa có danh mục thì không phát hành được.'", bold_prefix="Mâu thuẫn giữa Báo cáo và Mã nguồn: ")
    add_bullet(doc, "Thực tế tại src/LCMS.Application/RateCards/Commands/ComposeTariffCommand.cs dòng 295: Hệ thống vẫn hardcode ChargeCode = ReferenceCharge.Freight ('FREIGHT') cho mọi quy tắc ma trận! Đồng thời tại PublishRateVersionCommandHandler.cs dòng 239: Hệ thống chủ động bỏ qua kiểm tra danh mục đối với quy tắc ma trận (!rule.Matrix).", bold_prefix="Code thực tế vẫn hardcode: ")
    add_bullet(doc, "Hệ quả: Một bảng giá Advanced mới vẫn tự động được gán FREIGHT và phát hành thành công bình thường ngay cả khi tenant có 0 danh mục khoản giá chuẩn. Tuyên bố của Dev AI là sai lệch thực tế đối với Bảng giá nâng cao.", bold_prefix="Rủi ro nghiệp vụ: ")

    # Lỗ hổng 4
    add_heading_styled(doc, "Lỗ hổng 4: Import Bảng giá thiếu hoàn toàn kiểm soát WeightBand", level=3)
    add_bullet(doc, "Trong ImportRateCardsCommand.cs, đối tượng RateImportRule hoàn toàn không có các thuộc tính WeightBasis, WeightMin, WeightMax cho các quy tắc cơ bản (chỉ có Breaks). Quá trình ImportPreview và Commit không thực thi kiểm tra tính hợp lệ của bộ ba trọng lượng theo chuẩn RATE-UI-06.", bold_prefix="Bỏ lọt kiểm soát dữ liệu nạp: ")

    # Lỗ hổng 5
    add_heading_styled(doc, "Lỗ hổng 5: Bằng chứng UAT thiếu trung thực về kiểm thử giao diện Responsive", level=3)
    add_bullet(doc, "File capture.mjs phục vụ chụp ảnh bằng chứng được viết cứng độ rộng màn hình width: 1920 (Desktop tiêu chuẩn).", bold_prefix="Chỉ chụp Desktop Full HD: ")
    add_bullet(doc, "Không có bất kỳ bằng chứng kiểm thử nào trên độ phân giải máy tính bảng (768px), điện thoại (375px), hoặc kiểm thử thanh cuộn ngang khi bảng quy tắc có thêm cột Điều kiện trọng lượng và Khoản giá. Dev vẫn tự ý tick 'Hoàn thành' mục UI responsive.", bold_prefix="Thiếu bằng chứng xác thực: ")

    # Lỗ hổng 6
    add_heading_styled(doc, "Lỗ hổng 6: Trải nghiệm người dùng (UX) rơi vào 'Ngõ cụt' khi gặp GAP DATA-01", level=3)
    add_bullet(doc, "Khi tenant chưa có danh mục Khoản mục kinh tế, form hiển thị 'Chưa có danh mục khoản giá' và cảnh báo vàng.", bold_prefix="Không có lối thoát cho người dùng: ")
    add_bullet(doc, "Giao diện không cung cấp bất kỳ liên kết điều hướng (link/button) hay chỉ dẫn nào để nhân viên giá chuyển sang màn hình Quản lý Khoản mục kinh tế khai báo. Người dùng bị bế tắc hoàn toàn trong luồng thao tác mà không hiểu phải làm gì tiếp theo.", bold_prefix="Thiếu chỉ dẫn khắc phục: ")

    # Lỗ hổng 7
    add_heading_styled(doc, "Lỗ hổng 7: Kịch bản kiểm thử tự động mang tính 'chọn bài dễ để pass'", level=3)
    add_bullet(doc, "Trong bài test Uat02, dev chỉ kiểm tra chiều ngược (header BUY + rule customer), không kiểm tra chiều gửi cùng header qua AddPricingRule xem có bị lưu dư thừa party_type_code hay không.", bold_prefix="Kiểm thử phiến diện: ")
    add_bullet(doc, "Không có bài test nào kiểm tra kịch bản chuyển đổi qua lại giữa Bảng giá cơ bản và Nâng cao (chuyển mode làm mất hoặc bảo toàn dữ liệu khoản giá/trọng lượng).", bold_prefix="Bỏ qua kịch bản chuyển trạng thái: ")

    # Bảng tổng hợp đánh giá
    add_heading_styled(doc, "2.3. Bảng đối chiếu hiện trạng và đánh giá sơ bộ", level=2)
    tbl_eval = doc.add_table(rows=1, cols=4)
    format_table(tbl_eval, [1.5, 2.5, 1.3, 1.7],
                 ["Hạng mục", "Hiện trạng triển khai", "Dev tự chấm", "Thẩm định sơ bộ"],
                 [
                     ["RATE-UI-05 (Loại giá)", "Web tốt; API AddPricingRule vẫn nhận và lưu partyTypeCode dư thừa.", "PASS", "ĐẠT CÓ ĐIỀU KIỆN (80%)"],
                     ["RATE-UI-06 (Trọng lượng)", "SaveRateCard tốt; API AddPricingRule & Import hổng hoàn toàn validation.", "PASS", "KHÔNG ĐẠT (60%)"],
                     ["RATE-UI-07 (Khoản giá)", "Web basic tốt; Advanced vẫn hardcode FREIGHT và bypass kiểm tra master.", "PASS", "KHÔNG ĐẠT (65%)"],
                     ["Mobile App Integration", "apps/mobile/rates.tsx chưa hề được cập nhật, tạo rule sẽ bị kẹt publish.", "Hoàn thành", "LIỆT (0%) - BỎ QUÊN"],
                     ["Import Validation", "Không hỗ trợ nấc trọng lượng chuẩn trong RateImportRule.", "Hoàn thành", "CHƯA ĐẠT (40%)"],
                     ["Responsive UI", "Chỉ có ảnh 1920px, không có bằng chứng tablet/mobile.", "Hoàn thành", "THIẾU BẰNG CHỨNG (50%)"],
                     ["Chất lượng Test Suite", "414 test pass nhưng test kịch bản hẹp, bỏ qua tích hợp chéo.", "PASS", "TRUNG BÌNH (70%)"]
                 ])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # PHẦN 3
    # -------------------------------------------------------------
    add_heading_styled(doc, "PHẦN 3: PROMPT CHỈ THỊ YÊU CẦU DEV TÁI CHẤM ĐIỂM KHẮT KHE NHẤT", level=1)
    
    add_p(doc, "Dưới đây là nội dung Prompt chỉ thị toàn diện, đanh thép và chuẩn mực được thiết kế để PO / Tech Lead copy gửi trực tiếp cho Dev AI (hoặc đội ngũ lập trình). Prompt này thiết lập thang điểm 100 điểm với chế tài trừ điểm không khoan nhượng:")

    add_callout(doc,
                "HƯỚNG DẪN SỬ DỤNG CHO PO / USER:\n"
                "Hãy sao chép toàn bộ nội dung trong khung viền dưới đây (từ mục [BẮT ĐẦU PROMPT] đến [KẾT THÚC PROMPT]) và dán vào khung chat để ra lệnh cho Dev AI tiến hành tự thẩm định và khắc phục.",
                title="HƯỚNG DẪN VẬN HÀNH",
                callout_type="info")

    prompt_text = (
"""[BẮT ĐẦU PROMPT - CHỈ THỊ TÁI THẨM ĐỊNH & CHẤM ĐIỂM LẠI KHẮT KHE RATE-UI-05..07]

Gửi Dev AI phụ trách triển khai gói RATE-UI-05 → RATE-UI-07:

Báo cáo bàn giao tại commit bfd3d9e và tài liệu LCMS_UAT_RATE_UI_05_07_DEV_Result_v1.0.docx tuyên bố '100% PASS, 20/20 UAT, 0 skipped' ĐÃ BỊ BAN THẨM ĐỊNH VÀ PO BÁC BỎ do phát hiện nhiều lỗ hổng phòng thủ server-side, bỏ quên nền tảng di động và làm đẹp số liệu kiểm thử.

Bạn được yêu cầu TIẾN HÀNH TÁI CHẤM ĐIỂM LẠI CÔNG VIỆC CỦA CHÍNH MÌNH theo thang điểm 100 điểm dưới đây với tinh thần KHẮT KHE NHẤT, KHÔNG KHOAN NHƯỢNG, VÀ BẮT BUỘC KHẮC PHỤC TRIỆT ĐỂ:

══════════════════════════════════════════════════════════════════════════════════
A. THANG ĐIỂM 100 ĐIỂM - TIÊU CHÍ ĐÁNH GIÁ KHẮT KHE (STRICT AUDIT RUBRIC)
══════════════════════════════════════════════════════════════════════════════════

1. TRỤ CỘT 1: TÍNH ĐÚNG ĐẮN NGHIỆP VỤ & ĐỘ PHỦ ĐẶC TẢ (Tối đa 30 điểm)
- [10đ] Header party_type là nguồn chân lý duy nhất. Quy tắc kế thừa hoàn toàn. Bỏ dropdown chiều giá ở mọi nơi.
- [10đ] Điều kiện trọng lượng (Không lọc nấc / Gross / Chargeable): khóa/xóa nấc khi không lọc; validate Từ ≥ 0, Đến > Từ; đúng ngữ nghĩa nửa khoảng [Từ, Đến) cơ bản và inclusive nâng cao.
- [10đ] Nhóm cạnh tranh (charge_code): ẩn khỏi UI, derive từ danh mục master economic_charge_types. Chặn phát hành nếu thiếu master. Tuyệt đối không tự ý gán FREIGHT cho bất kỳ loại bảng giá nào (kể cả Advanced).

2. TRỤ CỘT 2: TÍNH TOÀN VẸN HỢP ĐỒNG API & PHÒNG THỦ SERVER-SIDE (Tối đa 25 điểm)
- [10đ] Nhất quán validation trên TẤT CẢ các đường ghi: Không chỉ POST/PUT /api/rate-cards mà cả POST /api/rate-versions/{id}/rules (AddPricingRuleCommand) đều phải chặn lỗi bộ ba trọng lượng (WeightBand.DeclarationError) và xung đột partyTypeCode.
- [10đ] Hỗ trợ đầy đủ EconomicChargeTypeId trên API thêm quy tắc đơn lẻ. Không chấp nhận API con để lọt dữ liệu rác mà màn hình chính chặn.
- [5đ] Ngăn chặn bypass: Không cho phép bảng giá Advanced tự ý vượt rào publish khi danh mục master chưa được định nghĩa.

3. TRỤ CỘT 3: ĐỒNG BỘ ĐA NỀN TẢNG (MOBILE, IMPORT & RESPONSIVE) (Tối đa 20 điểm)
- [10đ] Ứng dụng Di động (apps/mobile/app/modules/rates.tsx): Phải tương thích với chính sách mới. Nếu cho tạo quy tắc từ mobile thì phải có Khoản giá hoặc chặn có kiểm soát, không để tạo ra quy tắc 'chết' không thể publish.
- [5đ] Import Bảng giá: Hỗ trợ và validate chuẩn xác các trường trọng lượng và khoản giá chuẩn trong RateImportRule.
- [5đ] Responsive UI: Đảm bảo giao diện bảng khai báo 15 cột hiển thị chuẩn xác, không vỡ layout trên tablet (768px) và mobile (375px).

4. TRỤ CỘT 4: CHẤT LƯỢNG KIỂM THỬ & BẰNG CHỨNG THỰC NGHIỆM (Tối đa 15 điểm)
- [5đ] Kiểm thử biên: Test đầy đủ mốc 0, cận biên [Từ, Đến), chuỗi số thập phân.
- [5đ] Kiểm thử chéo (Integration Test): Test luồng AddPricingRule trực tiếp, test luồng Mobile, test chuyển đổi chế độ Basic <-> Advanced.
- [5đ] Bằng chứng xác thực: Chụp ảnh thực tế trên nhiều kích thước màn hình (responsive viewports), không dùng script chụp cứng 1920px.

5. TRỤ CỘT 5: TÍNH BẤT BIẾN, TƯƠNG THÍCH NGƯỢC & AN TOÀN DỮ LIỆU (Tối đa 10 điểm)
- [5đ] Dữ liệu lịch sử đã published không bị sửa đổi, snapshot pricing_rules và rating_details.rule_code giữ nguyên.
- [5đ] Không phát sinh migration thừa, có phương án rollback sạch (git revert an toàn).

══════════════════════════════════════════════════════════════════════════════════
B. DANH MỤC TRỪ ĐIỂM THẲNG TAY (AUTOMATIC PENALTIES)
══════════════════════════════════════════════════════════════════════════════════
- TRỪ 10 ĐIỂM: Bỏ quên màn hình Mobile apps/mobile/rates.tsx nhưng vẫn tự đánh dấu hoàn thành 100%.
- TRỪ 8 ĐIỂM: API AddPricingRuleCommandHandler thiếu kiểm tra WeightBand.DeclarationError, tạo lỗ hổng backend.
- TRỪ 7 ĐIỂM: Bảng giá Advanced vẫn hardcode FREIGHT và bypass kiểm tra master trong PublishRateVersionCommandHandler, trái ngược với báo cáo.
- TRỪ 5 ĐIỂM: Import Bảng giá không validate bộ ba trọng lượng.
- TRỪ 5 ĐIỂM: Bằng chứng giao diện chỉ có ảnh 1920px, thiếu bằng chứng responsive tablet/mobile.
- TRỪ 5 ĐIỂM: UX gặp ngõ cụt khi thiếu danh mục, không có link/button điều hướng giải quyết.

══════════════════════════════════════════════════════════════════════════════════
C. YÊU CẦU ĐẦU RA BẮT BUỘC CHO DEV AI
══════════════════════════════════════════════════════════════════════════════════
1. Bảng tự chấm điểm lại (Self-Audit Scorecard) chi tiết từng tiêu chí kèm số điểm bị trừ thực tế.
2. Bản phân tích nguyên nhân gốc rễ (Root Cause Analysis) của 7 lỗ hổng nêu trên.
3. Mã nguồn sửa đổi (Code Diffs) giải quyết triệt để các lỗ hổng:
   - Bổ sung WeightBand.DeclarationError vào AddPricingRuleCommandHandler.
   - Chuẩn hóa việc lưu party_type_code = null trên AddPricingRuleCommandHandler.
   - Xử lý đồng bộ cho apps/mobile/rates.tsx.
   - Làm rõ cơ chế và ràng buộc Khoản giá đối với Bảng giá Advanced.
   - Bổ sung link điều hướng UX khi gặp GAP DATA-01.
4. Bổ sung test tự động kiểm thử các kịch bản hổng và chạy lại toàn bộ test suite.

[KẾT THÚC PROMPT]"""
    )
    
    add_callout(doc, prompt_text, title="PROMPT LỆNH TÁI THẨM ĐỊNH CHO DEV AI", callout_type="prompt")

    # -------------------------------------------------------------
    # PHẦN 4
    # -------------------------------------------------------------
    add_heading_styled(doc, "PHẦN 4: HƯỚNG DẪN HÀNH ĐỘNG DÀNH CHO PO & TECH LEAD", level=1)
    
    add_p(doc, "Để đảm bảo an toàn tuyệt đối trước khi release lên môi trường Production cms-sg-01, Ban thẩm định khuyến nghị PO thực hiện quy trình kiểm soát 3 bước sau:")
    
    add_bullet(doc, "Gửi nguyên văn nội dung Prompt tại Phần 3 vào phiên làm việc của Dev AI, yêu cầu AI giải trình và gửi lại bản vá mã nguồn kèm bảng tự chấm điểm.", "Bước 1: Ban hành chỉ thị tái thẩm định: ")
    add_bullet(doc, "Yêu cầu Dev AI sửa mã nguồn theo 4 ưu tiên cấp bách: (1) Khóa lỗ hổng AddPricingRuleCommand; (2) Sửa logic hardcode FREIGHT trên Advanced nếu PO yêu cầu bắt buộc có master; (3) Đồng bộ module Mobile; (4) Thêm link điều hướng UX trên màn hình Web khi gặp GAP DATA-01.", "Bước 2: Yêu cầu khắc phục mã nguồn (Remediation): ")
    add_bullet(doc, "Chỉ ký phê duyệt bàn giao (Sign-off) khi: Điểm tự chấm sau khắc phục đạt tối thiểu 95/100 điểm, không còn điểm liệt ở các lỗ hổng API/Mobile, và có đủ bằng chứng test tích hợp thực tế.", "Bước 3: Nghiệm thu có điều kiện (Conditional Sign-off): ")

    # Save document
    out_dir = "docs/po/17102026-LCMS_UAT_RATE_UI_05_07_DEV_Package_v1.0"
    out_path = os.path.join(out_dir, "LCMS_UAT_RATE_UI_05_07_DEV_Audit_and_Strict_Grading_Prompt_v1.0.docx")
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    main()

