# -*- coding: utf-8 -*-
"""PO/BA confirmation — D03 relationships, TD2 integration, UI feasibility."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(r"c:\A1\git\cms\docs\reports\CMS_XacNhan_D03_TD2_UI_2026-09-18.docx")


def shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def run(p, text, *, bold=False, size=11, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def para(doc, text, *, bold=False, size=11):
    p = doc.add_paragraph()
    run(p, text, bold=bold, size=size)
    p.paragraph_format.space_after = Pt(6)
    return p


def table(doc, headers, rows, hdr_color="1F4E79"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.paragraphs[0].clear()
        run(cell.paragraphs[0], h, bold=True, size=9, color=(255, 255, 255))
        shade(cell, hdr_color)
    for ri, row in enumerate(rows, start=1):
        for ci, v in enumerate(row):
            cell = t.rows[ri].cells[ci]
            cell.paragraphs[0].clear()
            run(cell.paragraphs[0], str(v), size=8)
        if ri % 2 == 0:
            for ci in range(len(headers)):
                shade(t.rows[ri].cells[ci], "F5F5F5")
    doc.add_paragraph()


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.5)
        s.left_margin = s.right_margin = Cm(1.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(
        title,
        "CMS / LCMS — Xác nhận D03 Relationship Matrix, TD2 Integration & UI feasibility",
        bold=True,
        size=14,
    )

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(
        sub,
        f"Ngày: {date.today().isoformat()}  ·  Trả lời 3 điểm PO/BA trước khi chốt phạm vi implementation  ·  "
        "Chưa mở rộng sang TMS",
        size=9,
    )

    doc.add_heading("Bối cảnh đã thống nhất", level=1)
    para(
        doc,
        "Operational System = Operational Source of Truth (SoT). LCMS không trở thành TMS. "
        "D03 Operational Reference vẫn gồm Order / Bill / Shipment / Leg / Movement + relationships.",
    )
    para(
        doc,
        "Implementation hiện tại: Bill = đầy đủ; Order / Shipment / Leg / Movement = thin backend/API "
        "(C-002 external identity); UI và production Ops Sync chưa hoàn chỉnh. Team xác nhận đúng.",
    )

    doc.add_heading("1. Relationship Matrix thực tế", level=1)
    para(
        doc,
        "Tiêu chí: IMPLEMENTED = domain + bảng/FK + lệnh link/upsert + API đọc/ghi cạnh (graph hoặc route). "
        "PARTIAL = thiếu lớp. MISSING = chưa có quan hệ bền vững.",
    )
    table(
        doc,
        ["Quan hệ", "Cardinality / bảng", "Status", "Ghi chú"],
        [
            (
                "Order ↔ Bill",
                "N:N order_bill_links",
                "IMPLEMENTED",
                "Link API + GET /api/bills/{id}/graph",
            ),
            (
                "Bill ↔ Shipment",
                "N:N bill_shipment_links",
                "IMPLEMENTED",
                "Link API + graph",
            ),
            (
                "Shipment → Leg",
                "1:N transport_legs.shipment_id",
                "IMPLEMENTED",
                "FK bắt buộc lúc upsert leg",
            ),
            (
                "Bill ↔ Leg",
                "N:N bill_leg_links",
                "IMPLEMENTED",
                "Direct + union qua shipment trên graph",
            ),
            (
                "Leg ↔ Movement",
                "N:N leg_movement_links",
                "IMPLEMENTED",
                "Link API + graph",
            ),
            (
                "Bill ↔ Movement",
                "N:N bill_movement_links",
                "IMPLEMENTED",
                "Direct + union qua leg trên graph",
            ),
        ],
    )
    para(
        doc,
        "Kết luận điểm 1: Không quan hệ nào MISSING hoặc PARTIAL ở lớp persistence/API. "
        "Gap chung (không đổi status matrix): chưa có unlink; entity vẫn thin; UI Bill drawer mới hiện "
        "Order/Shipment — chưa render Leg/Movement.",
        bold=True,
    )

    doc.add_heading("2. Integration — TD2 baseline vs khi có Ops cụ thể", level=1)
    table(
        doc,
        ["Thành phần", "Bắt buộc TD2 baseline (chưa nối Ops)", "Khi có Operational System cụ thể"],
        [
            (
                "E03 push upsert/link (PUT Order/Shipment/Leg/Movement + link + bill graph)",
                "BẮT BUỘC — đã có",
                "Adapter đẩy vào các API này",
            ),
            (
                "C-002 idempotency trên operational refs",
                "BẮT BUỘC",
                "Giữ nguyên",
            ),
            (
                "D12 integration_records / integration_errors (+ outbox stub)",
                "BẮT BUỘC skeleton — đã có",
                "Pipeline thật ghi vào đây",
            ),
            (
                "Extract / vendor client",
                "Không",
                "BẮT BUỘC",
            ),
            (
                "Webhook ingress",
                "Không (chưa có code)",
                "BẮT BUỘC nếu Ops push event",
            ),
            (
                "Ops orchestrator / Integration Worker map payload → command",
                "Không",
                "BẮT BUỘC",
            ),
            (
                "Broker + outbox publish đầy đủ",
                "Không (stub OK)",
                "Cần khi sync async tin cậy",
            ),
        ],
    )
    para(
        doc,
        "Kết luận điểm 2: TD2 baseline = hợp đồng nhận/map operational refs + idempotency/traceability. "
        "Không yêu cầu extract / webhook / orchestrator sống để hoàn thành baseline. "
        "Connector Ops = dự án adapter sau khi PO chọn Ops SoT (xem docs/implementation-plan.md).",
        bold=True,
    )

    doc.add_heading(
        "3. UI — LIST / SEARCH / DETAIL / RELATIONSHIP / CROSS-NAV (không CRUD; trước Ops prod)",
        level=1,
    )
    para(
        doc,
        "Tạm không bàn CRUD Order/Shipment (H-002). Đánh giá khả thi trên dữ liệu/API hiện có.",
    )
    table(
        doc,
        ["Capability", "Order", "Shipment", "Leg", "Movement"],
        [
            (
                "LIST",
                "Ready (GET /api/orders)",
                "Cần thin List API",
                "Cần thin List API",
                "Cần thin List API",
            ),
            (
                "SEARCH",
                "Partial (search trả Bill hit theo order)",
                "Cần mở rộng search",
                "Cần mở rộng search",
                "Cần mở rộng search",
            ),
            (
                "DETAIL",
                "Ready (GET /api/orders/{id})",
                "Cần thin Get",
                "Cần thin Get",
                "Cần thin Get",
            ),
            (
                "RELATIONSHIP VIEW",
                "Bill-centric Ready (graph)",
                "Bill-centric Ready",
                "Bill-centric Ready",
                "Bill-centric Ready (flat; leg ids chưa expose)",
            ),
            (
                "CROSS-NAVIGATION",
                "Cần include billIds + trang/BFF",
                "Cần Get + related ids",
                "Cần Get (ShipmentId đã có trên entity)",
                "Cần expose link Leg↔Movement",
            ),
        ],
    )
    para(
        doc,
        "Kết luận điểm 3: Khả thi trước Ops production — Có. Seed bằng upsert/link API (tests Sprint 2 đã làm). "
        "Không blocked bởi Ops Sync.",
        bold=True,
    )
    para(doc, "Nếu PO muốn lát cắt nhỏ nhất để “nhìn được graph”:")
    para(doc, "1) Bill drawer — hiện thêm Leg/Movement + deep-link.")
    para(doc, "2) Thin List/Get cho Shipment / Leg / Movement (+ demo seed).")
    para(doc, "3) Search đa entity (đã ghi residual go-live).")

    doc.add_heading("Tóm tắt cho quyết định phạm vi PO/BA", level=1)
    para(
        doc,
        "Matrix quan hệ đã đủ backend. Thiếu chủ yếu: read API entity-centric + UI Bill graph (Leg/Movement) "
        "+ (tùy chọn) demo seed. Connector Ops chưa bắt buộc cho TD2 baseline. Chưa yêu cầu mở rộng sang TMS.",
        bold=True,
    )

    doc.add_heading("Tham chiếu kỹ thuật", level=1)
    para(
        doc,
        "API: src/LCMS.Api/Endpoints/OperationalReferenceEndpoints.cs · "
        "Graph: GetBillGraphQuery · Link entities trong LCMS.Domain · "
        "Tests: Sprint2OperationalReferenceTests / Sprint2FullOperationalReferenceTests · "
        "D12: IntegrationRecord + AuditIntegrationEndpoints · "
        "UI: BillFinancialDrawer (tab Liên quan) · "
        "Handoff: docs/handoff.md (2026-09-18).",
        size=9,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
