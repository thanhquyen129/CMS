from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

out = Path(r"c:\A1\git\cms\docs\po\LCMS_RATE_CR_20261014_02_v1.1_FINAL_POLISH.docx")
img = Path(r"c:\A1\git\cms\docs\po\rate-ac-evidence\polish")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

doc.add_heading("LCMS-RATE-CR-20261014-02 v1.1 — Hoàn thiện giao diện", 0)
p = doc.add_paragraph("Ngày: 14/10/2026. Phạm vi: form khai báo bảng giá trên web. Backend, migration và 36 AC đã PASS từ trước; tài liệu này chỉ ghi 4 hạng mục hiển thị.")

doc.add_heading("1. Kết quả", 1)
for line in [
    "Thanh Hủy / Lưu nháp không còn nổi đè lên khối hình thức khai báo và bảng quy tắc. Thanh này nằm cuối form.",
    "Badge đỏ “1 Issue” hết. Nguyên nhân là hydration: mã dòng quy tắc dùng số ngẫu nhiên nên id trên server khác id trên client. Mã dòng đầu và khung ma trận mặc định giờ cố định.",
    "Đã có 3 tuyến đang hoạt động SGN → HAN, SGN → DAD, HAN → DAD và 2 hãng Vietnam Airlines, Vietjet Air trên môi trường chụp.",
    "Cột Tuyến, Hãng vận chuyển, Nấc trọng lượng, Cách tính và Số tiền / Đơn giá có bề rộng tối thiểu. Nhãn “Đơn giá (theo số lượng cước)” không bị cắt ở 1920×1080. Ở 1366×768 bảng cuộn ngang; các cột tuyến, hãng, nấc và cách tính nằm trong khung nhìn.",
    "Zoom 125% tại 1366×768: không có ô nhập nào giao với thanh Hủy / Lưu nháp.",
]:
    doc.add_paragraph(line, style="List Bullet")

doc.add_heading("2. Ảnh nghiệm thu", 1)

shots = [
    ("basic-1366x768.png", "Ảnh 1a — Bảng giá cơ bản, 1366×768. Tuyến SGN → HAN, hãng Vietnam Airlines, loại trọng lượng thực, từ 0 kg đến 45 kg, cách tính đơn giá × số lượng, đơn vị theo số lượng cước. Thanh Hủy / Lưu nháp nằm dưới bảng, không che ô nhập."),
    ("basic-1920x1080.png", "Ảnh 1b — Bảng giá cơ bản, 1920×1080. Thêm cột Số tiền / Đơn giá: 120000 và nhãn Đơn giá (theo số lượng cước) đủ chữ. Thanh cuộn ngang vẫn còn cho các cột phía sau."),
    ("advanced-1366x768.png", "Ảnh 2a — Bảng giá nâng cao, 1366×768. Nhóm 2 HAN → DAD, hãng Vietjet Air, nấc [0, 45) đơn giá 31 và 41. Phía trên là nấc của nhóm 1."),
    ("advanced-1920x1080.png", "Ảnh 2b — Bảng giá nâng cao, 1920×1080. Nhóm 1 SGN → HAN / Vietnam Airlines, đơn giá 11 và 21. Nhóm 2 HAN → DAD / Vietjet Air, đơn giá 31 và 41."),
    ("detail-1366x768.png", "Ảnh 3a — Chi tiết bảng giá đã lưu, 1366×768, chỉ đọc. Mã POLISH-SGN-HAN, quy tắc FREIGHT, đơn giá 120.000 đ, lọc SGN-HAN · Vietnam Airlines · trọng lượng thực < 45."),
    ("detail-1920x1080.png", "Ảnh 3b — Cùng màn chi tiết, 1920×1080."),
]
for name, caption in shots:
    doc.add_paragraph(caption)
    doc.add_picture(str(img / name), width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("3. Cổng bàn giao", 1)
for line in [
    "Không còn badge đỏ 1 Issue trên form tạo bảng giá.",
    "Không còn thanh Hủy / Lưu nháp đè lên radio hay ô nhập.",
    "Ảnh trong tài liệu này thay bộ ảnh cũ bị đè nút, badge đỏ và danh mục tuyến trống.",
]:
    doc.add_paragraph(line, style="List Bullet")

doc.save(out)
print(out, out.stat().st_size)
